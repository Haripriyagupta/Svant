"""
Document and source code text extraction engine for SVANT.
Extracts content safely from TXT, MD, CSV, PDF, DOCX, Code, and Config files.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import NamedTuple, Optional

from svant.config import settings
from svant.core.classifier import FileClassifier, FileCategory
from svant.logger import get_logger

logger = get_logger("svant.core.extractor")


class ExtractionResult(NamedTuple):
    file_path: str
    status: str  # 'extracted', 'empty', 'skipped', 'failed'
    text: Optional[str]
    char_count: int
    word_count: int
    error: Optional[str] = None


class DocumentExtractor:
    """Extracts searchable plain text from various file formats."""

    def __init__(self, max_size_mb: Optional[int] = None) -> None:
        self.max_size_bytes = (max_size_mb or settings.max_extract_size_mb) * 1024 * 1024

    def extract(self, file_path: Path | str) -> ExtractionResult:
        p = Path(file_path)
        if not p.is_file():
            return ExtractionResult(
                file_path=str(p),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error="File does not exist or is not a regular file",
            )

        try:
            file_size = p.stat().st_size
        except OSError as e:
            return ExtractionResult(
                file_path=str(p),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error=f"Permission or OS error accessing file: {e}",
            )

        if file_size == 0:
            return ExtractionResult(
                file_path=str(p),
                status="empty",
                text="",
                char_count=0,
                word_count=0,
            )

        if file_size > self.max_size_bytes:
            return ExtractionResult(
                file_path=str(p),
                status="skipped",
                text=None,
                char_count=0,
                word_count=0,
                error=f"File exceeds maximum extraction size limit ({self.max_size_bytes // (1024 * 1024)} MB)",
            )

        classification = FileClassifier.classify(p)
        if not classification.is_text_extractable:
            return ExtractionResult(
                file_path=str(p),
                status="skipped",
                text=None,
                char_count=0,
                word_count=0,
                error=f"Binary or non-extractable file category ({classification.category.value})",
            )

        ext = p.suffix.lower()

        try:
            if ext == ".pdf":
                return self._extract_pdf(p)
            elif ext == ".docx":
                return self._extract_docx(p)
            elif ext in {".csv", ".tsv"}:
                return self._extract_csv(p, delimiter="," if ext == ".csv" else "\t")
            else:
                return self._extract_plain_text(p)
        except Exception as e:
            logger.warning(f"Failed to extract text from {p}: {e}")
            return ExtractionResult(
                file_path=str(p),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error=str(e),
            )

    def _extract_plain_text(self, path: Path) -> ExtractionResult:
        """Decode plain text or code using fallback encodings."""
        raw_bytes = path.read_bytes()
        # Binary check: null byte heuristic
        if b"\x00" in raw_bytes[:4096]:
            return ExtractionResult(
                file_path=str(path),
                status="skipped",
                text=None,
                char_count=0,
                word_count=0,
                error="File appears to be binary",
            )

        text = None
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = raw_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            return ExtractionResult(
                file_path=str(path),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error="Failed to decode text with supported character encodings",
            )

        stripped = text.strip()
        words = len(stripped.split()) if stripped else 0
        return ExtractionResult(
            file_path=str(path),
            status="extracted" if stripped else "empty",
            text=text,
            char_count=len(text),
            word_count=words,
        )

    def _extract_csv(self, path: Path, delimiter: str = ",") -> ExtractionResult:
        """Extract CSV text and format structured rows."""
        text_result = self._extract_plain_text(path)
        if text_result.status != "extracted" or not text_result.text:
            return text_result

        try:
            reader = csv.reader(io.StringIO(text_result.text), delimiter=delimiter)
            lines = []
            row_count = 0
            for row in reader:
                if row:
                    lines.append(" | ".join(row))
                    row_count += 1
                if row_count >= 1000:  # Cap at first 1000 rows to prevent memory explosion
                    lines.append("... [CSV preview truncated for indexing]")
                    break
            joined = "\n".join(lines)
            return ExtractionResult(
                file_path=str(path),
                status="extracted",
                text=joined,
                char_count=len(joined),
                word_count=len(joined.split()),
            )
        except Exception:
            # Fall back to the raw plain text extraction
            return text_result

    def _extract_pdf(self, path: Path) -> ExtractionResult:
        """Extract text from PDF pages using pypdf."""
        try:
            import pypdf
        except ImportError:
            return ExtractionResult(
                file_path=str(path),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error="pypdf package is not installed",
            )

        try:
            reader = pypdf.PdfReader(str(path))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return ExtractionResult(
                        file_path=str(path),
                        status="failed",
                        text=None,
                        char_count=0,
                        word_count=0,
                        error="Password-protected PDF cannot be extracted",
                    )

            pages_text = []
            for i, page in enumerate(reader.pages):
                try:
                    pt = page.extract_text()
                    if pt:
                        pages_text.append(pt)
                except Exception as page_err:
                    logger.debug(f"Could not extract page {i} of {path}: {page_err}")

            full_text = "\n\n".join(pages_text).strip()
            return ExtractionResult(
                file_path=str(path),
                status="extracted" if full_text else "empty",
                text=full_text,
                char_count=len(full_text),
                word_count=len(full_text.split()) if full_text else 0,
            )
        except Exception as e:
            return ExtractionResult(
                file_path=str(path),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error=f"PDF extraction error: {e}",
            )

    def _extract_docx(self, path: Path) -> ExtractionResult:
        """Extract text from DOCX using python-docx."""
        try:
            import docx
        except ImportError:
            return ExtractionResult(
                file_path=str(path),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error="python-docx package is not installed",
            )

        try:
            doc = docx.Document(str(path))
            elements = []
            for para in doc.paragraphs:
                if para.text.strip():
                    elements.append(para.text.strip())

            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        elements.append(" | ".join(row_texts))

            full_text = "\n\n".join(elements).strip()
            return ExtractionResult(
                file_path=str(path),
                status="extracted" if full_text else "empty",
                text=full_text,
                char_count=len(full_text),
                word_count=len(full_text.split()) if full_text else 0,
            )
        except Exception as e:
            return ExtractionResult(
                file_path=str(path),
                status="failed",
                text=None,
                char_count=0,
                word_count=0,
                error=f"DOCX extraction error: {e}",
            )
