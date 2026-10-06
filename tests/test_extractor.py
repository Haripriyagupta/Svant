"""
Tests for Document and Code Extraction Service.
"""

from pathlib import Path
import docx
from pypdf import PdfWriter
from svant.core.extractor import DocumentExtractor


def test_extract_txt_and_md(tmp_path: Path):
    txt_file = tmp_path / "hello.txt"
    txt_file.write_text("Hello SVANT World!\nThis is a plain text file.", encoding="utf-8")

    md_file = tmp_path / "readme.md"
    md_file.write_text("# SVANT Architecture\nMarkdown content here.", encoding="utf-8")

    extractor = DocumentExtractor()
    res_txt = extractor.extract(txt_file)
    assert res_txt.status == "extracted"
    assert "Hello SVANT World!" in res_txt.text
    assert res_txt.word_count > 0

    res_md = extractor.extract(md_file)
    assert res_md.status == "extracted"
    assert "SVANT Architecture" in res_md.text


def test_extract_csv(tmp_path: Path):
    csv_file = tmp_path / "data.csv"
    csv_file.write_text("name,role,level\nAlice,Developer,Senior\nBob,Security,Lead\n", encoding="utf-8")

    extractor = DocumentExtractor()
    res = extractor.extract(csv_file)
    assert res.status == "extracted"
    assert "Alice | Developer | Senior" in res.text


def test_extract_docx(tmp_path: Path):
    docx_file = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_paragraph("This is a paragraph inside a DOCX document.")
    table = doc.add_table(rows=1, cols=2)
    row_cells = table.rows[0].cells
    row_cells[0].text = "Header 1"
    row_cells[1].text = "Header 2"
    doc.save(str(docx_file))

    extractor = DocumentExtractor()
    res = extractor.extract(docx_file)
    assert res.status == "extracted"
    assert "This is a paragraph inside a DOCX document." in res.text
    assert "Header 1 | Header 2" in res.text


def test_extract_pdf(tmp_path: Path):
    # Create a minimal valid PDF using pypdf
    pdf_file = tmp_path / "test.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with open(pdf_file, "wb") as f:
        writer.write(f)

    extractor = DocumentExtractor()
    res = extractor.extract(pdf_file)
    # Blank page returns empty status safely without error
    assert res.status in ("empty", "extracted")


def test_extract_edge_cases(tmp_path: Path):
    extractor = DocumentExtractor(max_size_mb=1)

    # 1. Non-existent file
    res_missing = extractor.extract(tmp_path / "missing.txt")
    assert res_missing.status == "failed"

    # 2. Empty file
    empty_file = tmp_path / "empty.txt"
    empty_file.write_text("", encoding="utf-8")
    res_empty = extractor.extract(empty_file)
    assert res_empty.status == "empty"

    # 3. Binary file
    bin_file = tmp_path / "binary.bin"
    bin_file.write_bytes(b"\x00\x01\x02\x03\x00\x05")
    res_bin = extractor.extract(bin_file)
    assert res_bin.status == "skipped"
